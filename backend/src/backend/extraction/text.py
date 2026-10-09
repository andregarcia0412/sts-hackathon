"""Raw text of each uploaded file. No interpretation here: the agent decides what each file is."""

import hashlib
import io
import json
from dataclasses import dataclass, field
from typing import Any, Literal

from openpyxl import load_workbook
from pypdf import PdfReader

RawFormat = Literal["pdf", "table", "json", "text", "binary"]
TEXT_SUFFIXES = (".md", ".txt", ".markdown")


@dataclass
class RawFile:
    path: str
    sha256: str
    format: RawFormat
    lines: list[str]
    line_pages: list[int] | None = None
    json_data: Any = None
    error: str | None = None
    extra: dict[str, Any] = field(default_factory=dict)

    @property
    def text(self) -> str:
        return "\n".join(self.lines)


def _decode(data: bytes) -> str | None:
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        if encoding == "latin-1" and any(ord(ch) < 9 for ch in text[:2000]):
            return None
        return text
    return None


def _pdf(path: str, digest: str, data: bytes) -> RawFile:
    try:
        reader = PdfReader(io.BytesIO(data))
        lines: list[str] = []
        pages: list[int] = []
        for number, page in enumerate(reader.pages, start=1):
            page_lines = (page.extract_text() or "").splitlines()
            lines += page_lines
            pages += [number] * len(page_lines)
        return RawFile(path, digest, "pdf", lines, line_pages=pages)
    except Exception as error:  # corrupted PDF: no text, recorded
        return RawFile(path, digest, "pdf", [], line_pages=[], error=f"pdf: {type(error).__name__}")


def _xlsx(path: str, digest: str, data: bytes) -> RawFile:
    try:
        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        lines = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if value is None else str(value) for value in row]
            while cells and cells[-1] == "":
                cells.pop()
            lines.append(";".join(cells))
        while lines and lines[-1] == "":
            lines.pop()
        return RawFile(path, digest, "table", lines)
    except Exception as error:
        return RawFile(path, digest, "binary", [], error=f"xlsx: {type(error).__name__}")


def read_raw(path: str, data: bytes) -> RawFile:
    digest = hashlib.sha256(data).hexdigest()
    lower = path.lower()
    if lower.endswith(".pdf") or data.startswith(b"%PDF"):
        return _pdf(path, digest, data)
    if lower.endswith((".xlsx", ".xlsm")) or data.startswith(b"PK\x03\x04"):
        return _xlsx(path, digest, data)
    text = _decode(data)
    if text is None:
        return RawFile(path, digest, "binary", [])
    lines = text.splitlines()
    if lower.endswith(".json"):
        try:
            return RawFile(path, digest, "json", lines, json_data=json.loads(text))
        except json.JSONDecodeError as error:
            return RawFile(path, digest, "json", lines, error=f"json: {error.msg}")
    if lower.endswith(".csv"):
        return RawFile(path, digest, "table", lines)
    return RawFile(path, digest, "text", lines)
